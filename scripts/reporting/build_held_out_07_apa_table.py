from pathlib import Path
import csv

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "results/phase_learning/final_lopo/held_out_07_report/held_out_07_3s_result_table.csv"
OUTPUT = ROOT / "results/phase_learning/final_lopo/held_out_07_report/Held_Out_07_LOPO_APA_Table.docx"


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        element = tc_mar.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def clear_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn("w:" + name))
        if edge is None:
            edge = OxmlElement("w:" + name)
            borders.append(edge)
        edge.set(qn("w:val"), "nil")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, attributes in edges.items():
        edge = borders.find(qn("w:" + edge_name))
        if edge is None:
            edge = OxmlElement("w:" + edge_name)
            borders.append(edge)
        for key, value in attributes.items():
            edge.set(qn("w:" + key), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def style_run(run, size=9.5, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_cell_text(cell, text, *, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.2):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


with SOURCE.open("r", encoding="utf-8-sig", newline="") as stream:
    rows = list(csv.DictReader(stream))

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Cm(1.15)
section.bottom_margin = Cm(1.15)
section.left_margin = Cm(1.45)
section.right_margin = Cm(1.45)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(9.5)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
run = p.add_run("Table 1")
style_run(run, size=10.5, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
run = p.add_run("Held-Out Participant 07 LOPO Classification and Transition Results Within a ±3-s Tolerance")
style_run(run, size=10.5, italic=True)

table = doc.add_table(rows=2 + len(rows), cols=8)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
clear_table_borders(table)

widths_cm = [4.0, 2.55, 2.55, 2.75, 2.55, 2.55, 2.55, 2.3]
for row in table.rows:
    for cell, width in zip(row.cells, widths_cm):
        cell.width = Cm(width)

model_header = table.cell(0, 0).merge(table.cell(1, 0))
overall_header = table.cell(0, 1).merge(table.cell(0, 2))
phase_header = table.cell(0, 3).merge(table.cell(0, 6))
transition_header = table.cell(0, 7).merge(table.cell(1, 7))

set_cell_text(model_header, "Model", bold=True)
set_cell_text(overall_header, "Overall", bold=True)
set_cell_text(phase_header, "Phase-specific F1", bold=True)
set_cell_text(transition_header, "Transition F1", bold=True)

subheaders = ["Accuracy", "Macro F1", "Free Motion", "Approach", "Painting", "Retract"]
for col, label in enumerate(subheaders, start=1):
    set_cell_text(table.cell(1, col), label, bold=True, size=8.8)

rule = {"val": "single", "sz": "8", "color": "000000"}
for cell in table.rows[0].cells:
    set_cell_border(cell, top=rule)
for cell in table.rows[1].cells:
    set_cell_border(cell, bottom=rule)

fields = [
    "Accuracy_3s_pct",
    "Macro_F1_3s_pct",
    "FREE_F1_3s_pct",
    "APPROACH_F1_3s_pct",
    "PAINTING_F1_3s_pct",
    "RETRACT_F1_3s_pct",
    "Transition_F1_3s_pct",
]

for row_index, item in enumerate(rows, start=2):
    set_cell_text(table.cell(row_index, 0), item["Model"], align=WD_ALIGN_PARAGRAPH.LEFT)
    for col, field in enumerate(fields, start=1):
        set_cell_text(table.cell(row_index, col), f'{float(item[field]):.1f}')

for cell in table.rows[-1].cells:
    set_cell_border(cell, bottom=rule)

set_repeat_table_header(table.rows[0])
set_repeat_table_header(table.rows[1])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(7)
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.0
r = p.add_run("Note. ")
style_run(r, size=8.7, italic=True)
r = p.add_run(
    "Values are percentages for held-out participant 07. Classification and phase-specific F1 scores "
    "use the previously applied same-phase ±3-s temporal-tolerance criterion. Transition F1 matches "
    "the nearest unused event with the same from-to phase direction within ±3 s. LOPO = leave-one-participant-out."
)
style_run(r, size=8.7)

core = doc.core_properties
core.title = "Held-Out Participant 07 LOPO Results"
core.subject = "APA-style result table"
core.author = ""
core.last_modified_by = ""

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
