import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "results" / "phase_learning" / "final_lopo" / "results" / "tolerant_3s" / "phase_metrics_3s.json"
WORK = ROOT / ".cache" / "document_previews" / "lopo_word_assets"
OUTPUT = ROOT / "results" / "phase_learning" / "final_lopo" / "results" / "tolerant_3s" / "LOPO_Results_APA_Figures_Table_and_Timeline.docx"
WORK.mkdir(parents=True, exist_ok=True)

FONT_REGULAR = Path(r"C:\Windows\Fonts\arial.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\arialbd.ttf")


def font(size, bold=False):
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size=size)


def draw_centered(draw, xy, text, fnt, fill, anchor="mm"):
    draw.text(xy, text, font=fnt, fill=fill, anchor=anchor)


def draw_grouped_bar_chart(
    output_path,
    title,
    models,
    series,
    colors,
    width=3000,
    height=1700,
):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    left, right = 245, 110
    top, bottom = 255, 245
    plot_left, plot_right = left, width - right
    plot_top, plot_bottom = top, height - bottom
    plot_width = plot_right - plot_left
    plot_height = plot_bottom - plot_top
    y_max = 110.0

    title_font = font(54, bold=True)
    axis_font = font(34)
    tick_font = font(29)
    label_font = font(27, bold=True)
    legend_font = font(30)

    draw_centered(draw, (width / 2, 68), title, title_font, "#222222", anchor="ma")

    legend_y = 165
    item_widths = []
    for item in series:
        bbox = draw.textbbox((0, 0), item[0], font=legend_font)
        item_widths.append(44 + 16 + (bbox[2] - bbox[0]) + 55)
    legend_start = (width - sum(item_widths)) / 2
    cursor = legend_start
    for (name, _, _), color, item_w in zip(series, colors, item_widths):
        draw.rectangle((cursor, legend_y - 16, cursor + 44, legend_y + 18), fill=color)
        draw.text((cursor + 60, legend_y + 2), name, font=legend_font, fill="#333333", anchor="lm")
        cursor += item_w

    def y_px(value):
        return plot_bottom - (value / y_max) * plot_height

    for tick in range(0, 101, 20):
        y = y_px(tick)
        draw.line((plot_left, y, plot_right, y), fill="#D9DDE2", width=2)
        draw.text((plot_left - 24, y), str(tick), font=tick_font, fill="#333333", anchor="rm")

    draw.line((plot_left, plot_top, plot_left, plot_bottom), fill="#333333", width=4)
    draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill="#333333", width=4)
    # Rotate the y-axis label after drawing it onto a transparent tile.
    label_box = Image.new("RGBA", (420, 70), (255, 255, 255, 0))
    label_draw = ImageDraw.Draw(label_box)
    label_draw.text((210, 35), "Score (%)", font=axis_font, fill="#222222", anchor="mm")
    label_box = label_box.rotate(90, expand=True)
    image.paste(label_box, (35, int((height - label_box.height) / 2)), label_box)
    draw = ImageDraw.Draw(image)

    n_groups = len(models)
    n_series = len(series)
    group_w = plot_width / n_groups
    cluster_ratio = 0.76 if n_series == 4 else 0.58
    cluster_w = group_w * cluster_ratio
    gap = 15 if n_series == 4 else 28
    bar_w = (cluster_w - gap * (n_series - 1)) / n_series

    for group_idx, model in enumerate(models):
        group_center = plot_left + group_w * (group_idx + 0.5)
        cluster_left = group_center - cluster_w / 2
        for series_idx, ((_, means, sds), color) in enumerate(zip(series, colors)):
            mean = means[group_idx]
            sd = sds[group_idx]
            x0 = cluster_left + series_idx * (bar_w + gap)
            x1 = x0 + bar_w
            y_mean = y_px(mean)
            draw.rectangle((x0, y_mean, x1, plot_bottom), fill=color)

            x_center = (x0 + x1) / 2
            upper = min(y_max, mean + sd)
            lower = max(0.0, mean - sd)
            y_upper = y_px(upper)
            y_lower = y_px(lower)
            cap = min(32, bar_w * 0.35)
            draw.line((x_center, y_upper, x_center, y_lower), fill="#202020", width=5)
            draw.line((x_center - cap, y_upper, x_center + cap, y_upper), fill="#202020", width=5)
            draw.line((x_center - cap, y_lower, x_center + cap, y_lower), fill="#202020", width=5)

            # Put the mean label above the upper error-bar cap, with a consistent gap.
            label_y = y_px(min(y_max - 1.0, upper + 2.4))
            draw.text((x_center, label_y), f"{mean:.1f}", font=label_font, fill="#202020", anchor="ms")

        label_y = plot_bottom + 42
        lines = model.split("\n")
        for line_idx, line in enumerate(lines):
            draw.text(
                (group_center, label_y + line_idx * 36),
                line,
                font=tick_font,
                fill="#333333",
                anchor="ma",
            )

    image.save(output_path, dpi=(300, 300), optimize=True)


def set_cell_text(cell, text, bold=False, italic=False, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        edge_data = kwargs.get(edge)
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ["val", "sz", "space", "color"]:
            if key in edge_data:
                element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=10, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_apa_figure_note(document, number, text):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Figure {number}. ")
    set_run_font(r, 9.5, bold=True)
    r = p.add_run(text)
    set_run_font(r, 9.5)


with SOURCE.open("r", encoding="utf-8") as f:
    data = json.load(f)

summary = data["lopo_summary"]
display_models = ["Causal\nFSM", "Random\nForest", "GMM", "Gaussian\nHMM", "GMM–\nHMM", "Causal\nLSTM"]


def pct(key):
    return [row[key] * 100.0 for row in summary]


phase_series = [
    ("Free Motion", pct("free_motion_f1_3s_mean"), pct("free_motion_f1_3s_sd")),
    ("Approach", pct("approach_f1_3s_mean"), pct("approach_f1_3s_sd")),
    ("Painting", pct("painting_f1_3s_mean"), pct("painting_f1_3s_sd")),
    ("Retract", pct("retract_f1_3s_mean"), pct("retract_f1_3s_sd")),
]
overall_series = [
    ("Accuracy", pct("tolerant_accuracy_3s_mean"), pct("tolerant_accuracy_3s_sd")),
    ("Macro F1", pct("tolerant_macro_f1_3s_mean"), pct("tolerant_macro_f1_3s_sd")),
]

phase_figure = WORK / "lopo_phase_specific_f1_revised.png"
overall_figure = WORK / "lopo_performance_revised.png"
timeline_figure = ROOT / "results" / "phase_learning" / "final_lopo" / "results" / "tolerant_3s" / "lopo_08_01_phase_timeline.png"
draw_grouped_bar_chart(
    phase_figure,
    "LOPO Phase-Specific F1",
    display_models,
    phase_series,
    ["#587EA3", "#D69A4B", "#A8D08D", "#B8646B"],
)
draw_grouped_bar_chart(
    overall_figure,
    "LOPO Performance",
    display_models,
    overall_series,
    ["#4F759B", "#D48A40"],
)

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11.69)
section.page_height = Inches(8.27)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(10)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(3)
r = title.add_run("LOPO Results")
set_run_font(r, 17, bold=True)

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.paragraph_format.space_after = Pt(6)
r = intro.add_run("Leave-one-participant-out performance across six held-out participants.")
set_run_font(r, 10)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(0)
r = p.add_run()
r.add_picture(str(phase_figure), width=Inches(9.75))
add_apa_figure_note(
    doc,
    1,
    "Phase-specific F1 scores. Bars represent means and error bars represent ±1 SD across six LOPO folds.",
)

doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(0)
r = p.add_run()
r.add_picture(str(overall_figure), width=Inches(9.75))
add_apa_figure_note(
    doc,
    2,
    "Overall LOPO performance. Bars represent means and error bars represent ±1 SD across six LOPO folds.",
)

doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(0)
r = p.add_run()
r.add_picture(str(timeline_figure), width=Inches(10.0))
add_apa_figure_note(
    doc,
    3,
    "Framewise causal phase predictions for trial 08_01. Participant 08 was excluded from training; gray intervals indicate gaps between valid segments.",
)

doc.add_page_break()
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Table 1")
set_run_font(r, 11, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("LOPO Classification and Transition Results Within a ±3-s Tolerance")
set_run_font(r, 11, italic=True)

table = doc.add_table(rows=2 + len(summary), cols=8)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

widths = [1.52, 1.12, 1.12, 1.22, 1.20, 1.18, 1.14, 1.08]
for row in table.rows:
    row.height = Inches(0.37)
    for idx, cell in enumerate(row.cells):
        cell.width = Inches(widths[idx])

model_header = table.cell(0, 0).merge(table.cell(1, 0))
overall_header = table.cell(0, 1).merge(table.cell(0, 2))
phase_header = table.cell(0, 3).merge(table.cell(0, 6))
transition_header = table.cell(0, 7)

set_cell_text(model_header, "Model", bold=True, size=8.5)
set_cell_text(overall_header, "Overall", bold=True, size=8.5)
set_cell_text(phase_header, "Phase-specific F1", bold=True, size=8.5)
set_cell_text(transition_header, "Transition", bold=True, size=8.5)

subheaders = ["Accuracy", "Macro F1", "Free Motion", "Approach", "Painting", "Retract", "F1"]
for col, text_value in enumerate(subheaders, start=1):
    set_cell_text(table.cell(1, col), text_value, bold=True, size=8.0)

row_names = ["Causal FSM", "Random Forest", "GMM", "Gaussian HMM", "GMM–HMM", "Causal LSTM"]
metrics = [
    ("tolerant_accuracy_3s_mean", "tolerant_accuracy_3s_sd"),
    ("tolerant_macro_f1_3s_mean", "tolerant_macro_f1_3s_sd"),
    ("free_motion_f1_3s_mean", "free_motion_f1_3s_sd"),
    ("approach_f1_3s_mean", "approach_f1_3s_sd"),
    ("painting_f1_3s_mean", "painting_f1_3s_sd"),
    ("retract_f1_3s_mean", "retract_f1_3s_sd"),
    ("transition_f1_3s_mean", "transition_f1_3s_sd"),
]

for row_idx, (row_data, row_name) in enumerate(zip(summary, row_names), start=2):
    set_cell_text(table.cell(row_idx, 0), row_name, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    for col_idx, (mean_key, sd_key) in enumerate(metrics, start=1):
        mean = row_data[mean_key] * 100.0
        sd = row_data[sd_key] * 100.0
        set_cell_text(table.cell(row_idx, col_idx), f"{mean:.1f} ({sd:.1f})", size=8.5)

# APA-style minimal horizontal rules: top, header separator, and bottom only.
rule = {"val": "single", "sz": "8", "space": "0", "color": "000000"}
none = {"val": "nil"}
for r_idx, row in enumerate(table.rows):
    for cell in row.cells:
        set_cell_border(cell, top=none, bottom=none, left=none, right=none)
        if r_idx == 0:
            set_cell_border(cell, top=rule)
        if r_idx == 1:
            set_cell_border(cell, bottom=rule)
        if r_idx == len(table.rows) - 1:
            set_cell_border(cell, bottom=rule)

set_repeat_table_header(table.rows[0])
set_repeat_table_header(table.rows[1])

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(5)
note.paragraph_format.space_after = Pt(0)
r = note.add_run("Note. ")
set_run_font(r, 9, italic=True)
r = note.add_run(
    "Values are percentages and are reported as M (SD) across six held-out participants. "
    "Classification and transition scores were calculated using a ±3-s tolerance. "
    "LOPO = leave-one-participant-out; M = mean; SD = standard deviation."
)
set_run_font(r, 9)

doc.core_properties.title = "LOPO Results: Figures and APA Table"
doc.core_properties.subject = "Phase recognition LOPO evaluation"
doc.core_properties.author = "OpenAI"
doc.save(OUTPUT)

print(OUTPUT)
print(phase_figure)
print(overall_figure)
