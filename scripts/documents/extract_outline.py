from docx import Document
from pathlib import Path

source = Path(r"D:\Downloads\Outline_Comparative Evaluation of Visual Interfaces_CA (1).docx")
output = Path(__file__).resolve().parents[2] / "docs" / "reviews" / "visual_interfaces" / "outline_text.txt"

doc = Document(source)
lines = []
for i, paragraph in enumerate(doc.paragraphs, start=1):
    text = paragraph.text.strip()
    if text:
        lines.append(f"P{i:04d}\t[{paragraph.style.name}]\t{text}")

for ti, table in enumerate(doc.tables, start=1):
    lines.append(f"\nTABLE {ti}")
    for ri, row in enumerate(table.rows, start=1):
        values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        lines.append(f"T{ti}R{ri}\t" + "\t".join(values))

output.parent.mkdir(parents=True, exist_ok=True)
output.write_text("\n".join(lines), encoding="utf-8")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} output={output}")
